import re
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


def reporting_year(value: Any, row_number: int) -> int:
    if isinstance(value, (date, datetime)):
        return value.year
    if isinstance(value, str):
        reporting_date = value.strip()
        for date_format in ("%Y-%m-%d", "%m/%d/%Y", "%Y%m%d"):
            try:
                return (
                    datetime.strptime(reporting_date, date_format)
                    .replace(tzinfo=UTC)
                    .year
                )
            except ValueError:
                continue
    raise ValueError(
        f"Invalid or missing Reporting Date in Excel row {row_number}: {value!r}"
    )


def filing_rows(workbook_bytes: bytes) -> list[tuple[int, str, int]]:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ValueError("SEC Excel export does not contain a worksheet.")
        headers: dict[str, int] = {}
        for index, cell in enumerate(worksheet[1]):
            if cell.value is not None:
                header = re.sub(r"[^a-z]", "", str(cell.value).lower())
                headers[header] = index

        url_column = next(
            (
                headers[header]
                for header in ("filingsurl", "filingurl")
                if header in headers
            ),
            None,
        )
        date_column = headers.get("reportingdate")
        if url_column is None or date_column is None:
            raise ValueError(
                "SEC Excel export must contain 'Filings URL' and "
                "'Reporting Date' columns."
            )

        rows = []
        for row_number, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
            url_cell = row[url_column]
            url = url_cell.hyperlink.target if url_cell.hyperlink else url_cell.value
            if url is None or not str(url).strip():
                continue
            year = reporting_year(row[date_column].value, row_number)
            rows.append((year, str(url).strip(), row_number))
        return rows
    finally:
        workbook.close()
